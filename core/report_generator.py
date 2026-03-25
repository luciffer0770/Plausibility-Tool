"""Annotated Excel and PDF report generation."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from fpdf import FPDF
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models import Project, Status
from core.plausibility_engine import limits_display_str
from core.style_constants import STATUS_FAIL, STATUS_NO_DATA, STATUS_OK, STATUS_WARNING
from database.db_manager import DatabaseManager

logger = logging.getLogger(__name__)

_STATUS_FILL = {
    Status.OK.value: STATUS_OK,
    Status.WARNING.value: STATUS_WARNING,
    Status.FAIL.value: STATUS_FAIL,
    Status.NO_DATA.value: STATUS_NO_DATA,
}


def _fill(hex_color: str) -> PatternFill:
    hx = hex_color.lstrip("#")
    return PatternFill(start_color=hx, end_color=hx, fill_type="solid")


def write_annotated_excel(
    source_path: Path,
    output_path: Path,
    session_id: int,
    db: DatabaseManager,
) -> Path:
    """
    Write original data (sheet Original) plus Plausibility_Analysis sheet with results.

    Analysis sheet includes Status, Deviation %, Root cause with row fills.
    Original sheet is not modified.
    """
    source_path = Path(source_path)
    output_path = Path(output_path)
    suffix = source_path.suffix.lower()
    if suffix == ".csv":
        df_orig = pd.read_csv(source_path)
    elif suffix == ".xlsx":
        df_orig = pd.read_excel(source_path, engine="openpyxl")
    else:
        df_orig = pd.read_excel(source_path)

    meas = db.get_measurements_for_session(session_id)
    rows_out: list[dict[str, Any]] = []
    for m in meas:
        lo = m.get("limit_lower")
        hi = m.get("limit_upper")
        lim_s = ""
        if lo is not None or hi is not None:
            lim_s = f"{lo if lo is not None else ''} - {hi if hi is not None else ''}"
        rows_out.append(
            {
                "Parameter": m.get("parameter_name"),
                "Value": m.get("measured_value"),
                "Min": m.get("value_min"),
                "Max": m.get("value_max"),
                "Avg": m.get("value_avg"),
                "Lower_Limit": lo,
                "Upper_Limit": hi,
                "Limits": lim_s.strip(),
                "Status": m.get("status"),
                "Deviation_%": m.get("deviation"),
                "Root_Cause": m.get("root_cause") or "",
                "Corrective_Action": m.get("corrective_action") or "",
                "Timestamp": m.get("timestamp") or "",
            }
        )
    df_res = pd.DataFrame(rows_out)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_orig.to_excel(writer, index=False, sheet_name="Original")
        df_res.to_excel(writer, index=False, sheet_name="Plausibility_Analysis")
        ws = writer.sheets["Plausibility_Analysis"]
        status_col = list(df_res.columns).index("Status") + 1
        fills = {k: _fill(v) for k, v in _STATUS_FILL.items()}
        for ri in range(2, len(df_res) + 2):
            st = ws.cell(row=ri, column=status_col).value
            fl = fills.get(str(st))
            if fl:
                for ci in range(1, len(df_res.columns) + 1):
                    ws.cell(row=ri, column=ci).fill = fl
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wo = writer.sheets["Original"]
        wo.freeze_panes = "A2"
        wo.auto_filter.ref = wo.dimensions

    logger.info("Wrote annotated Excel: %s", output_path)
    return output_path


def write_plausibility_report_excel(
    output_path: Path,
    session_id: int,
    db: DatabaseManager,
) -> Path:
    """
    Bosch-style plausibility report (v3 §8): green header row, summary, data.
    """
    from openpyxl import Workbook

    output_path = Path(output_path)
    meas = db.get_measurements_for_session(session_id)
    sess = db.get_upload_session(session_id) or {}
    fname = sess.get("file_name", "")
    datum = sess.get("datum") or ""
    ver = sess.get("version_test") or ""
    nrows = sess.get("record_count") or 0

    ok_n = sum(1 for m in meas if m.get("status") == "OK")
    high_n = int(sess.get("above_count") or sum(1 for m in meas if m.get("status") == "HIGH"))
    low_n = int(sess.get("below_count") or sum(1 for m in meas if m.get("status") == "LOW"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    green = "006B3F"
    white = "FFFFFF"
    red = "ED0007"
    ok_green = "00884B"
    fail_fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")

    ws.merge_cells("A1:K1")
    c1 = ws["A1"]
    c1.value = "BOSCH Plausibility Check Report"
    c1.font = Font(bold=True, size=14, color=white)
    c1.fill = PatternFill(start_color=green, end_color=green, fill_type="solid")
    c1.alignment = Alignment(horizontal="left", vertical="center")

    ws["A2"] = f"File: {fname} | Date: {datum} | Version: {ver} | Runs: {nrows}"

    ws["A3"], ws["B3"] = "Total", len(meas)
    ws["B3"].font = Font(bold=True)
    ws["C3"], ws["D3"] = "Passed", ok_n
    ws["D3"].font = Font(bold=True, color=ok_green)
    ws["E3"], ws["F3"] = "Above", high_n
    ws["F3"].font = Font(bold=True, color=red if high_n else "000000")
    ws["G3"], ws["H3"] = "Below Low", low_n
    ws["H3"].font = Font(bold=True, color=red if low_n else "000000")

    headers = [
        "Parameter",
        "Description",
        "Category",
        "Type",
        "Runs",
        "Min",
        "Max",
        "Avg",
        "Limits",
        "Status",
        "Root Cause",
    ]
    hr = 5
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=hr, column=col, value=h)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill(start_color=green, end_color=green, fill_type="solid")

    sorted_meas = sorted(meas, key=lambda m: str(m.get("parameter_name", "")))
    for i, m in enumerate(sorted_meas, start=hr + 1):
        lo = m.get("limit_lower")
        hi = m.get("limit_upper")
        u = str(m.get("unit") or "")
        lims = limits_display_str(
            float(lo) if lo is not None else None,
            float(hi) if hi is not None else None,
            u,
        )
        st = str(m.get("status", ""))
        row = [
            m.get("parameter_name"),
            m.get("description") or "",
            m.get("category") or "",
            m.get("param_type") or "",
            m.get("num_runs"),
            m.get("value_min"),
            m.get("value_max"),
            m.get("value_avg"),
            lims,
            st,
            m.get("root_cause") or "",
        ]
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=col, value=val)
            if col in (6, 7, 8) and val is not None and isinstance(val, (int, float)):
                cell.number_format = "0.00"
            if st in ("HIGH", "LOW"):
                cell.fill = fail_fill
            if col == 10:
                if st == "OK":
                    cell.font = Font(bold=True, color=ok_green)
                elif st in ("HIGH", "LOW", "FAIL"):
                    cell.font = Font(bold=True, color=red)
                elif st in ("NO_DATA", "N/A"):
                    cell.font = Font(color="808080")

    ws.freeze_panes = f"A{hr + 1}"
    ws.auto_filter.ref = f"A{hr}:K{hr + len(sorted_meas)}"

    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 14

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    logger.info("Wrote plausibility report Excel: %s", output_path)
    return output_path


class _ReportPdf(FPDF):
    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}", align="C")


def write_pdf_summary(
    output_path: Path,
    project: Project,
    session_summary: dict[str, Any],
    fail_warn_rows: list[dict[str, Any]],
    logo_path: Optional[Path] = None,
) -> Path:
    """
    PDF with project block, counts, and FAIL/WARN table.

    Args:
        output_path: Destination .pdf path.
        project: Current project.
        session_summary: Keys ok_n, warn_n, fail_n, nd_n, file_name, upload_date.
        fail_warn_rows: Measurement dicts for non-OK.
        logo_path: Optional Bosch logo image.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = _ReportPdf()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    if logo_path and Path(logo_path).is_file():
        try:
            pdf.image(str(logo_path), x=10, y=8, w=28)
        except Exception as e:
            logger.warning("Could not embed logo: %s", e)

    pdf.set_y(36)
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Bosch Plausibility Check Report", ln=1)

    pdf.set_font("Helvetica", size=11)
    pdf.cell(0, 6, f"Project: {project.name}", ln=1)
    pdf.cell(0, 6, f"Engine: {project.engine_type_key()}", ln=1)
    pdf.cell(0, 6, f"Test bed: {project.test_bed_id or '-'}", ln=1)
    pdf.cell(0, 6, f"File: {session_summary.get('file_name', '-')}", ln=1)
    pdf.cell(0, 6, f"Upload: {session_summary.get('upload_date', '-')}", ln=1)
    pdf.ln(4)

    ok_n = int(session_summary.get("ok_n") or 0)
    w_n = int(session_summary.get("warn_n") or 0)
    f_n = int(session_summary.get("fail_n") or 0)
    nd_n = int(session_summary.get("nd_n") or 0)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Summary", ln=1)
    pdf.set_font("Helvetica", size=11)
    pdf.cell(
        0,
        6,
        f"OK: {ok_n}  |  WARNING: {w_n}  |  FAIL: {f_n}  |  NO DATA: {nd_n}",
        ln=1,
    )
    pdf.ln(6)

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 8, "FAIL and WARNING parameters", ln=1)
    pdf.set_font("Helvetica", size=9)
    col_w = [22, 22, 28, 50, 55]
    headers = ["Status", "Parameter", "Value", "Limits", "Root cause"]
    for i, h in enumerate(headers):
        pdf.cell(col_w[i], 7, h, border=1)
    pdf.ln(7)

    for m in fail_warn_rows:
        if m.get("status") not in ("FAIL", "WARNING", "HIGH", "LOW"):
            continue
        lo = m.get("limit_lower")
        hi = m.get("limit_upper")
        lim = "-"
        if lo is not None or hi is not None:
            lim = f"{lo if lo is not None else ''}-{hi if hi is not None else ''}"
        val = m.get("measured_value")
        val_s = f"{val:.3g}" if val is not None else "-"
        pdf.set_font("Helvetica", size=9)
        pdf.cell(col_w[0], 6, str(m.get("status", "")), border=1)
        pdf.cell(col_w[1], 6, str(m.get("parameter_name", "")), border=1)
        pdf.cell(col_w[2], 6, val_s, border=1)
        pdf.cell(col_w[3], 6, lim[:28], border=1)
        x = pdf.get_x()
        y = pdf.get_y()
        pdf.multi_cell(col_w[4], 6, str(m.get("root_cause", ""))[:200], border=1)
        pdf.set_xy(pdf.l_margin, y + 6)
        if pdf.get_y() > 270:
            pdf.add_page()

    pdf.output(str(output_path))
    logger.info("Wrote PDF: %s", output_path)
    return output_path
