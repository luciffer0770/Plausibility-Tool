"""Export failed (or selected) measurement rows to a small Excel file."""

from __future__ import annotations

from pathlib import Path
from typing import Any, List

from openpyxl import Workbook
from openpyxl.styles import Font

from core.plausibility_engine import limits_display_str


def export_measurements_excel(path: Path, rows: List[dict[str, Any]], title: str = "Results") -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    headers = (
        "Parameter",
        "Description",
        "Category",
        "Type",
        "Runs",
        "Min",
        "Max",
        "Avg",
        "Limits",
        "Status",
        "Root cause",
    )
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)

    def fmt(v: Any) -> Any:
        if v is None:
            return ""
        return v

    for ri, m in enumerate(rows, start=2):
        lo = m.get("limit_lower")
        hi = m.get("limit_upper")
        u = str(m.get("unit") or "")
        lims = limits_display_str(
            float(lo) if lo is not None else None,
            float(hi) if hi is not None else None,
            u,
        )
        ws.cell(row=ri, column=1, value=m.get("parameter_name"))
        ws.cell(row=ri, column=2, value=m.get("description") or "")
        ws.cell(row=ri, column=3, value=m.get("category") or "")
        ws.cell(row=ri, column=4, value=m.get("param_type") or m.get("parameter_type") or "")
        ws.cell(row=ri, column=5, value=m.get("num_runs"))
        ws.cell(row=ri, column=6, value=fmt(m.get("value_min")))
        ws.cell(row=ri, column=7, value=fmt(m.get("value_max")))
        ws.cell(row=ri, column=8, value=fmt(m.get("value_avg")))
        ws.cell(row=ri, column=9, value=lims)
        ws.cell(row=ri, column=10, value=m.get("status"))
        ws.cell(row=ri, column=11, value=m.get("root_cause") or "")

    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
